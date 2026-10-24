#!/usr/bin/env python3
"""Create editable article manuscripts from versioned Markdown sources.

Requires python-docx. PDF export and layout review are separate release steps.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run

ROOT = Path(__file__).resolve().parents[1]
PUB = ROOT / "docs/publications"


def mrun(text):
    run = OxmlElement("m:r")
    val = OxmlElement("m:t")
    val.text = text
    run.append(val)
    return run


def sub(text, index):
    e = OxmlElement("m:sSub")
    base = OxmlElement("m:e")
    base.append(mrun(text))
    s = OxmlElement("m:sub")
    s.append(mrun(index))
    e.extend([base, s])
    return e


def frac(numerator, denominator):
    f = OxmlElement("m:f")
    n = OxmlElement("m:num")
    d = OxmlElement("m:den")
    for item in numerator:
        n.append(item)
    for item in denominator:
        d.append(item)
    f.extend([n, d])
    return f


def math_sequence(text):
    pattern = re.compile(r"(CFADS|FCF|DSCR|Tax|rbar|lambda|yhat|[A-Za-z])_(?:\(([^)]+)\)|([A-Za-z0-9]+))")
    seq, pos = [], 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            seq.append(mrun(text[pos:match.start()].replace(" - ", " − ").replace(" * ", " · ")))
        symbol = {"lambda": "λ", "rbar": "r̄", "yhat": "ŷ"}.get(match[1], match[1])
        seq.append(sub(symbol, (match[2] or match[3]).replace("-", "−")))
        pos = match.end()
    if pos < len(text):
        seq.append(mrun(text[pos:].replace(" - ", " − ").replace(" * ", " · ")))
    return seq


def add_equation(doc, formula):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    math = OxmlElement("m:oMath")
    if formula.startswith("DSCR_t ="):
        nodes = [sub("DSCR", "t"), mrun(" = "), frac([sub("CFADS", "t")], [sub("I", "t"), mrun(" + "), sub("P", "t")])]
    elif formula.startswith("MAE ="):
        nodes = [mrun("MAE = "), frac([mrun("1")], [mrun("N")]), mrun(" ∑ "), mrun("|"), sub("e", "i"), mrun("|")]
    elif formula.startswith("RMSE ="):
        rad = OxmlElement("m:rad")
        prop = OxmlElement("m:radPr")
        hide = OxmlElement("m:degHide"); hide.set(qn("m:val"), "1"); prop.append(hide)
        degree = OxmlElement("m:deg")
        inside = OxmlElement("m:e")
        inside.append(frac([mrun("1")], [mrun("N")]))
        inside.append(mrun(" ∑ "))
        power = OxmlElement("m:sSup")
        a = OxmlElement("m:e"); a.append(sub("e", "i"))
        b = OxmlElement("m:sup"); b.append(mrun("2"))
        power.extend([a, b]); inside.append(power)
        rad.extend([prop, degree, inside])
        nodes = [mrun("RMSE = "), rad]
    else:
        nodes = math_sequence(formula)
    for node in nodes:
        math.append(node)
    p._p.append(math)


def hyperlink(paragraph, label, url):
    link = OxmlElement("w:hyperlink")
    rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    col = OxmlElement("w:color"); col.set(qn("w:val"), "274f72"); props.append(col)
    under = OxmlElement("w:u"); under.set(qn("w:val"), "single"); props.append(under)
    run.append(props)
    txt = OxmlElement("w:t"); txt.text = label; run.append(txt)
    link.append(run); paragraph._p.append(link)


INLINE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def plain_run(paragraph, text, **kwargs):
    # Format simple variable subscripts in running prose without raw notation.
    variable = re.compile(r"\b(CFADS|FCF|DSCR|Tax|rbar|lambda|yhat|[A-Za-z])_(?:\(([^)]+)\)|([A-Za-z0-9]+))")
    pos = 0
    for match in variable.finditer(text):
        if match.start() > pos:
            r = paragraph.add_run(text[pos:match.start()]); apply_run(r, **kwargs)
        r = paragraph.add_run({"lambda": "λ", "rbar": "r̄", "yhat": "ŷ"}.get(match[1], match[1])); apply_run(r, **kwargs)
        r = paragraph.add_run(match[2] or match[3]); r.font.subscript = True; apply_run(r, **kwargs)
        pos = match.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:]); apply_run(r, **kwargs)


def apply_run(run, bold=False, code=False):
    run.bold = bold
    if code:
        run.font.name = "DejaVu Sans Mono"
        run.font.size = Pt(9)


def add_inline(paragraph, text):
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            plain_run(paragraph, text[pos:match.start()])
        token = match[0]
        if token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            hyperlink(paragraph, link[1], link[2])
        elif token.startswith("**"):
            plain_run(paragraph, token[2:-2], bold=True)
        else:
            plain_run(paragraph, token[1:-1], code=True)
        pos = match.end()
    plain_run(paragraph, text[pos:])


def add_table(doc, lines):
    rows = [[v.strip() for v in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", v) for v in r)]
    n = len(rows[0])
    table = doc.add_table(rows=1, cols=n)
    table.autofit = False
    total = 6.8
    if n == 5:
        widths = [1.30, 1.25, 1.52, 1.3, 1.43] if rows[0][0] == "Сценарий" else [1.08, 1.25, 1.52, 1.3, 1.65]
    elif n == 4:
        widths = [2.3, 1.8, 1.3, 1.4] if rows[0][0] in {"Regime", "Режим"} else [2.6, 1.35, 1.35, 1.5]
    elif n == 3:
        if any("results/" in cell for row in rows for cell in row):
            widths = [1.55, 2.8, 2.45]
        elif len(rows) > 6:
            widths = [2.2, 2.9, 1.7]
        else:
            widths = [3.3, 1.75, 1.75]
    else:
        widths = [2.25, 4.55] if n == 2 else [total/n]*n
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, values in enumerate(rows):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        for j, text in enumerate(values):
            cell = cells[j]; cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + edge); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D9D9D9"); borders.append(el)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for side, val in (("top", 90), ("bottom", 90), ("left", 95), ("right", 95)):
                el = OxmlElement("w:"+side); el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa"); margins.append(el)
            tcpr.append(margins)
            shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "DCE6EE" if i == 0 else ("F5F7F9" if i%2==0 else "FFFFFF")); tcpr.append(shd)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
            add_inline(p, text)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if re.fullmatch(r"[-−+0-9.,% x]+", text) else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(9.5); r.bold = i == 0; r.font.color.rgb = RGBColor(0, 0, 0)
        trpr = table.rows[i]._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit"); trpr.append(no_split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader"); trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build(source, target, language):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.78); section.bottom_margin = Inches(.75)
    section.left_margin = Inches(.85); section.right_margin = Inches(.85)
    family = "Liberation Serif" if source.name.startswith("ssrn_") else "Liberation Sans"
    normal = doc.styles["Normal"]
    normal.font.name = family; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for key, size in [("Title", 22), ("Heading 1", 14), ("Heading 2", 12)]:
        style = doc.styles[key]; style.font.name = family; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = key != "Title"
        style.paragraph_format.space_before = Pt(14 if key != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    # The bundled base template contains an accent border on title styles.
    # Academic manuscripts use typography and whitespace without title rules.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.core_properties.author = "Dmitry Kazakov"
    doc.core_properties.subject = "Infrastructure cash flow forecasting with a bounded agent harness"
    doc.core_properties.comments = ""
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    lines = source.read_text(encoding="utf-8").splitlines()
    i, code, language_tag = 0, False, ""
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            code = not code; language_tag = line[3:]; i += 1; continue
        if code:
            p = doc.add_paragraph()
            continues = i+1 < len(lines) and not lines[i+1].startswith("```")
            p.paragraph_format.space_after = Pt(0 if continues else 7)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = continues
            r = p.add_run(line); r.font.name = "DejaVu Sans Mono"; r.font.size = Pt(8.4)
            i += 1; continue
        if not line.strip():
            i += 1; continue
        if line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i]); i += 1
            add_table(doc, table_lines); continue
        image = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image:
            p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(source.parent / image[2]), width=Inches(6.8))
            following = next((x for x in lines[i+1:] if x.strip()), "")
            if not following.startswith(("**Figure ", "**Рисунок ")):
                cap = doc.add_paragraph(image[1]); cap.paragraph_format.space_after = Pt(9)
                for run in cap.runs: run.font.size=Pt(9); run.italic=True
            i += 1; continue
        if line.startswith("# "):
            title = line[2:]
            clean_title = title.replace("Cash-Flow", "Cash Flow").replace(":", "\n").replace("—", " ")
            p = doc.add_paragraph(clean_title, style="Title")
            doc.core_properties.title = title
            i += 1; continue
        if line.startswith("### ") or line.startswith("## "):
            level = 2 if line.startswith("### ") else 1
            heading = line[level+2:]
            heading = heading.replace(":", " ").replace("—", " ")
            doc.add_heading(heading, level=level)
            i += 1; continue
        equation = re.fullmatch(r"`([^`]+)`[.,]?", line)
        if equation and "=" in equation[1]:
            add_equation(doc, equation[1]); i += 1; continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_inline(p, line[2:]); i += 1; continue
        p = doc.add_paragraph()
        add_inline(p, line)
        if line.startswith(("**Table ", "**Таблица ")):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
        if line.startswith(("**Figure ", "**Рисунок ")):
            p.paragraph_format.keep_together = True
            for run in p.runs:
                run.font.size = Pt(9.5)
        if line in ("**Dmitry Kazakov**", "**Дмитрий Казаков**", "24 October 2026", "dmkkazakov@gmail.com"):
            p.paragraph_format.space_after = Pt(5)
        if line.startswith("[1]") or (line.startswith("[") and "] " in line[:6]):
            p.paragraph_format.space_after = Pt(5)
            for r in p.runs: r.font.size = Pt(9.5)
        if language == "ru" and re.match(r"^\d+\.\s", line):
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.0
            for run_element in p._p.xpath(".//w:r"):
                Run(run_element, p).font.size = Pt(9.5)
        i += 1
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target", choices=("all", "habr", "ssrn", "ssrn-ru"), default="all")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs/publications")
    args = parser.parse_args()
    manuscripts = {
        "habr": ("habr_article.md", "Infra_Cashflow_Habr_Dmitry_Kazakov.docx", "ru"),
        "ssrn": ("ssrn_paper.md", "Infra_Cashflow_SSRN_Dmitry_Kazakov.docx", "en"),
        "ssrn-ru": ("ssrn_paper_ru.md", "Infra_Cashflow_SSRN_Dmitry_Kazakov_RU.docx", "ru"),
    }
    for name, (source, output, language) in manuscripts.items():
        if args.target in ("all", name):
            build(PUB / source, args.output_dir / output, language)
